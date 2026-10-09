"""Resume file validation and text extraction (PDF and DOCX).

Uploaded files are untrusted: we check the declared type, the file signature (magic bytes),
and the size, and extract plain text only. Nothing in the file is ever executed or rendered
as HTML.
"""

import io
import re
import zipfile
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from ..errors import ValidationFailed

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ALLOWED = {".pdf": PDF_MIME, ".docx": DOCX_MIME}
MAX_PAGES = 10
MAX_TEXT_CHARS = 60_000
MIN_TEXT_CHARS = 200


@dataclass
class ParsedResume:
    text: str
    mime_type: str
    extension: str
    page_count: int | None


def detect_type(filename: str, data: bytes) -> tuple[str, str]:
    """Return (extension, mime) after checking both the extension and the file signature."""
    name = (filename or "").lower().strip()
    ext = "." + name.rsplit(".", 1)[-1] if "." in name else ""
    if ext not in ALLOWED:
        raise ValidationFailed("Only PDF and DOCX files are supported.", code="unsupported_file_type")

    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValidationFailed("This file does not look like a valid PDF.", code="invalid_file")
    if ext == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ValidationFailed("This file does not look like a valid DOCX document.", code="invalid_file")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
                if "word/document.xml" not in names:
                    raise ValidationFailed("This DOCX file is missing its document body.", code="invalid_file")
                # Guard against zip bombs: limit total uncompressed size.
                if sum(i.file_size for i in zf.infolist()) > 50 * 1024 * 1024:
                    raise ValidationFailed("This DOCX file is too large when uncompressed.", code="invalid_file")
        except zipfile.BadZipFile as exc:
            raise ValidationFailed("This DOCX file is corrupted.", code="invalid_file") from exc
    return ext, ALLOWED[ext]


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise ValidationFailed("Password-protected PDFs are not supported.", code="encrypted_pdf") from exc
        pages = reader.pages
        if len(pages) > MAX_PAGES:
            raise ValidationFailed(f"Resumes longer than {MAX_PAGES} pages are not supported.", code="too_many_pages")
        text = "\n".join((p.extract_text() or "") for p in pages)
        return text, len(pages)
    except PdfReadError as exc:
        raise ValidationFailed("We couldn't read this PDF. Try exporting it again.", code="invalid_file") from exc


def _extract_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailed("We couldn't read this DOCX file.", code="invalid_file") from exc
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_resume(filename: str, data: bytes, max_bytes: int) -> ParsedResume:
    if not data:
        raise ValidationFailed("The uploaded file is empty.", code="empty_file")
    if len(data) > max_bytes:
        raise ValidationFailed(f"Files must be {max_bytes // (1024 * 1024)} MB or smaller.", code="file_too_large")

    ext, mime = detect_type(filename, data)
    if ext == ".pdf":
        raw, pages = _extract_pdf(data)
    else:
        raw, pages = _extract_docx(data), None

    text = _clean(raw)
    if len(text) < MIN_TEXT_CHARS:
        raise ValidationFailed(
            "We could not extract enough text from this file. If it is a scanned image, please upload a text-based PDF or a DOCX file.",
            code="no_text",
        )
    return ParsedResume(text=text, mime_type=mime, extension=ext, page_count=pages)


# ---------------------------------------------------------------------------
# Data minimisation before sending text to the AI provider
# ---------------------------------------------------------------------------
_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
# Candidate phone numbers; only sequences with 10-13 digits are redacted, so date ranges
# like "2021-2025" and scores like "8.5/10" are left untouched.
_PHONE_CANDIDATE = re.compile(r"(?<![\w.])\+?\(?\d[\d\s().-]{8,18}\d(?![\w])")
_URL = re.compile(r"\b(?:https?://)?(?:www\.)?([a-z0-9-]+(?:\.[a-z0-9-]+)*\.[a-z]{2,})(/[^\s)]*)?", re.IGNORECASE)
_KNOWN_LINK_HOSTS = ("linkedin.com", "github.com", "gitlab.com", "leetcode.com", "kaggle.com", "behance.net")


def redact_for_ai(text: str) -> str:
    """Remove direct contact details (emails, phone numbers, profile URLs) that the AI does
    not need for analysis. Profile links are replaced by their host so the analysis can still
    note that, for example, a GitHub profile is present."""

    def url_sub(m: re.Match) -> str:
        host = m.group(1).lower()
        for known in _KNOWN_LINK_HOSTS:
            if host.endswith(known):
                return f"[{known} link]"
        return "[link]" if m.group(0).lower().startswith(("http", "www")) else m.group(0)

    def phone_sub(m: re.Match) -> str:
        digits = sum(ch.isdigit() for ch in m.group(0))
        return "[phone]" if 10 <= digits <= 13 else m.group(0)

    text = _EMAIL.sub("[email]", text)
    text = _URL.sub(url_sub, text)
    text = _PHONE_CANDIDATE.sub(phone_sub, text)
    return text
