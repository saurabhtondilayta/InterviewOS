import io
import zipfile

import pytest
from docx import Document

from app.errors import ValidationFailed
from app.services.resume_parser import parse_resume, redact_for_ai

MAX = 5 * 1024 * 1024

RESUME_LINES = [
    "Education: B.Tech Computer Science, 2023 - 2027, CGPA 8.5/10",
    "Skills: Python, Java, SQL, React, Node.js, Docker",
    "Projects: Built a URL shortener handling 10k requests/day using FastAPI and Redis.",
    "Internship: Backend intern - improved API latency by 30% by adding caching.",
    "Certifications: AWS Certified Cloud Practitioner",
]


def make_pdf(lines: list[str]) -> bytes:
    """Build a minimal, valid single-page PDF with real text (no external tools needed)."""
    content = "BT /F1 11 Tf 50 750 Td 14 TL " + " ".join(f"({line}) '" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx(lines: list[str]) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Languages"
    table.rows[0].cells[1].text = "English, Hindi"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parses_pdf():
    parsed = parse_resume("Resume.PDF", make_pdf(RESUME_LINES), MAX)
    assert parsed.mime_type == "application/pdf"
    assert parsed.page_count == 1
    assert "URL shortener" in parsed.text


def test_parses_docx_including_tables():
    parsed = parse_resume("cv.docx", make_docx(RESUME_LINES), MAX)
    assert parsed.extension == ".docx"
    assert "AWS Certified" in parsed.text
    assert "Languages | English, Hindi" in parsed.text


@pytest.mark.parametrize("name", ["resume.exe", "resume.doc", "resume.txt", "resume"])
def test_rejects_unsupported_extensions(name):
    with pytest.raises(ValidationFailed) as e:
        parse_resume(name, make_pdf(RESUME_LINES), MAX)
    assert e.value.code == "unsupported_file_type"


def test_rejects_spoofed_pdf():
    with pytest.raises(ValidationFailed) as e:
        parse_resume("resume.pdf", b"MZ\x90\x00 this is an executable" * 20, MAX)
    assert e.value.code == "invalid_file"


def test_rejects_zip_that_is_not_docx():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(ValidationFailed):
        parse_resume("resume.docx", buf.getvalue(), MAX)


def test_rejects_oversized_and_empty():
    with pytest.raises(ValidationFailed) as e:
        parse_resume("r.pdf", b"%PDF-" + b"0" * 2000, 1000)
    assert e.value.code == "file_too_large"
    with pytest.raises(ValidationFailed) as e:
        parse_resume("r.pdf", b"", MAX)
    assert e.value.code == "empty_file"


def test_rejects_documents_without_text():
    with pytest.raises(ValidationFailed) as e:
        parse_resume("r.pdf", make_pdf(["Hi"]), MAX)
    assert e.value.code == "no_text"


class TestRedaction:
    def test_removes_contact_details(self):
        text = "Riya | riya.k@example.com | +91 98765 43210 | linkedin.com/in/riya-k | https://riya.dev/portfolio"
        out = redact_for_ai(text)
        assert "example.com" not in out and "[email]" in out
        assert "98765" not in out and "[phone]" in out
        assert "riya-k" not in out and "[linkedin.com link]" in out
        assert "riya.dev" not in out and "[link]" in out

    def test_keeps_dates_scores_and_tech_names(self):
        text = "B.Tech 2021-2025, CGPA 8.5/10. Built with Node.js, React.js/Next.js and ASP.NET. Reduced latency 30%."
        assert redact_for_ai(text) == text

    def test_github_link_is_labelled(self):
        assert redact_for_ai("Code: github.com/someone/project") == "Code: [github.com link]"
